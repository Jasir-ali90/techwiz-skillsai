"""Hidden-document readiness: a document the system has never seen, and a
role created during judging, must flow through with no code change."""
import io

import pytest
from docx import Document as DocxDocument

pytestmark = pytest.mark.db


def hidden_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Nexora Labs Machine Learning Operations Standard", level=0)
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph("1.1 This document defines how models are trained and released at Nexora Labs.")
    doc.add_heading("2. Model Training", level=1)
    doc.add_paragraph("2.1 Machine Learning Engineers must register every training dataset in the model registry "
                      "before training begins.")
    doc.add_paragraph("2.2 Machine Learning Engineers must complete the Responsible AI Basics module within ten "
                      "calendar days of their joining date.")
    doc.add_paragraph("2.3 It is recommended that experiments are tracked with a run identifier.")
    doc.add_heading("3. Model Release", level=1)
    doc.add_paragraph("3.1 A model must not be promoted to production without a bias evaluation report.")
    doc.add_paragraph("3.2 Data Analysts may request read access to evaluation dashboards.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def test_unseen_document_and_new_role_flow_end_to_end(client, admin):
    departments = {d["code"]: d["id"] for d in (await client.get("/departments", headers=admin)).json()}
    role = await client.post("/roles", headers=admin, json={
        "code": "ML_ENG", "title": "Machine Learning Engineer", "department_id": departments["DATA"]})
    assert role.status_code == 201
    role_id = role.json()["id"]

    upload = await client.post("/documents", headers=admin, data={
        "document_code": "STD-MLOPS-030", "title": "Machine Learning Operations Standard",
        "document_type": "PROCESS_DOCUMENT", "version": "1", "effective_date": "2026-08-01",
        "department_id": departments["DATA"]},
        files={"file": ("mlops.docx", hidden_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert upload.status_code == 201, upload.text
    doc_id = upload.json()["document"]["id"]
    parsed = (await client.post(f"/documents/{doc_id}/parse", headers=admin)).json()
    assert parsed["parser"] == "python-docx" and parsed["chunks_created"] >= 5 and parsed["chunks_embedded"] >= 5

    extracted = (await client.post("/matrix/extract", headers=admin, json={"document_id": doc_id})).json()
    assert extracted["created"] >= 4
    assert extracted["rejected_by_reason"].get("informational", 0) >= 1
    mapped = (await client.post("/matrix/map-roles", headers=admin, json={})).json()
    assert "Machine Learning Engineer" in mapped["per_role"]

    matrix = (await client.get(f"/matrix/role/{role_id}", headers=admin)).json()
    statements = {r["statement"]: r for r in matrix["requirements"]}
    registry = next(r for s, r in statements.items() if "model registry" in s)
    assert registry["mapping_method"] == "EXPLICIT_ROLE_MENTION" and registry["is_mandatory"]
    basics = next(r for s, r in statements.items() if "Responsible AI Basics" in s)
    assert basics["deadline_days"] == 10 and basics["requirement_type"] == "MUST_COMPLETE"

    employee = await client.post("/employees", headers=admin, json={
        "employee_code": "EMP-900", "full_name": "Nadia Rahman", "job_role_id": role_id,
        "department_id": departments["DATA"], "experience_level": "INTERMEDIATE", "joining_date": "2026-09-01"})
    assert employee.status_code == 201
    plan = await client.post("/plans/generate", headers=admin, json={"employee_id": employee.json()["id"]})
    assert plan.status_code == 200, plan.text
    body = plan.json()
    cited = {x["source_chunk_code"] for m in body["modules"] for x in m["tasks"] + m["checklist"]}
    assert any(c.startswith("STD-MLOPS-030") for c in cited)
    assert body["validation"]["metrics"]["mandatory_requirement_coverage_score"] == 100.0
