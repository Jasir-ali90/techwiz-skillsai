"""Document upload validators, parsing, chunking, requirement extraction, role
mapping and prerequisites. Pure Python against the real sample documents."""
from datetime import date
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from src.core.app_config import load_yaml
from src.core.textutil import deadline_days, durations
from src.document_processing.chunker import chunk_document
from src.document_processing.docx_parser import parse_docx
from src.document_processing.pdf_parser import parse_pdf
from src.document_validation import validators as v
from src.models.enums import MappingMethod, Priority, RequirementType
from src.role_matrix.extractor import ChunkInput, RuleSet, extract_from_chunk
from src.role_matrix.prerequisites import ReqNode, detect
from src.role_matrix.role_mapper import RequirementInput, RoleInput, RoleMapper
from src.security.adversarial import scan_text

SAMPLES = Path("sample_documents")
RULES = RuleSet(load_yaml("requirement_rules"))


def chunks_of(name: str, code: str):
    return chunk_document(parse_pdf(str(SAMPLES / name)), code)


def extract(name: str, code: str, dtype: str):
    out = []
    for r in chunks_of(name, code):
        outcome = extract_from_chunk(ChunkInput(None, r["chunk_code"], r["section_id"], r["heading"], r["content"],
                                                None, code, dtype, "Doc", bool(scan_text(r["content"]))), RULES)
        out.append(outcome)
    return out


# ------------------------------------------------------------ document upload validators
def test_magic_bytes_reject_renamed_executable():
    report = v.ValidationReport()
    v.check_file_type("policy.pdf", b"MZ\x90\x00" + b"\x00" * 100, report)
    assert not report.ok and "not a valid PDF" in report.failures[0]["message"]


def test_disallowed_extension_rejected():
    report = v.ValidationReport()
    v.check_file_type("policy.exe", b"%PDF-1.7", report)
    assert not report.ok


def test_empty_oversized_dates_and_version():
    report = v.ValidationReport()
    v.check_not_empty(b"", report)
    v.check_file_size(b"x" * (2 * 1024 * 1024), 1, report)
    v.check_dates(date(2026, 5, 1), date(2026, 4, 1), report)
    v.check_version(2, 2, report)
    assert {c["check"] for c in report.failures} == {"empty_document", "file_size", "dates", "version"}


def test_valid_pdf_passes_type_check():
    report = v.ValidationReport()
    v.check_file_type("ok.pdf", (SAMPLES / "SOP-DEPLOY-007_v1.pdf").read_bytes(), report)
    assert report.ok


# ------------------------------------------------------------ parsing
def test_pdf_parser_keeps_page_numbers():
    parsed = parse_pdf(str(SAMPLES / "POL-INFOSEC-001_v1.pdf"))
    assert parsed.page_count >= 1 and parsed.blocks
    assert all(b.page_number for b in parsed.blocks)


def test_docx_parser_keeps_paragraph_index_and_headings(tmp_path):
    doc = DocxDocument()
    doc.add_heading("1. Leave Entitlement", level=1)
    doc.add_paragraph("1.1 All employees must submit leave requests within three calendar days of their joining date.")
    doc.add_paragraph("1.2 It is recommended that leave is planned a month ahead.")
    path = tmp_path / "leave.docx"
    doc.save(path)
    parsed = parse_docx(str(path))
    assert parsed.blocks[0].is_heading
    assert [b.paragraph_index for b in parsed.blocks] == [1, 2, 3]
    chunks = chunk_document(parsed, "POL-LEAVE-009")
    assert [c["section_id"] for c in chunks] == ["1.1", "1.2"]
    assert chunks[0]["heading_path"] == "Leave Entitlement"


# ------------------------------------------------------------ chunking
def test_chunks_follow_numbered_sections_with_codes_and_headings():
    chunks = chunks_of("POL-INFOSEC-001_v1.pdf", "POL-INFOSEC-001")
    by_section = {c["section_id"]: c for c in chunks if c["section_id"]}
    assert "2.3" in by_section and by_section["2.3"]["heading"] == "Account Security"
    assert by_section["2.3"]["chunk_code"].startswith("POL-INFOSEC-001#C")
    assert by_section["2.3"]["content"].startswith("2.3 Employees must complete")
    codes = [c["chunk_code"] for c in chunks]
    assert len(codes) == len(set(codes))


def test_faq_answers_inherit_their_question_section():
    chunks = chunks_of("FAQ-ENG-002_v1.pdf", "FAQ-ENG-002")
    answer = next(c for c in chunks if "full two weeks" in c["content"])
    assert answer["section_id"] == "2.1"


# ------------------------------------------------------------ requirement extraction
def test_infosec_yields_know_complete_and_recommended_and_rejects_purpose():
    outcomes = extract("POL-INFOSEC-001_v1.pdf", "POL-INFOSEC-001", "INFOSEC_POLICY")
    reqs = [r for o in outcomes for r in o.requirements]
    types = {r.requirement_type for r in reqs}
    assert {RequirementType.MUST_KNOW, RequirementType.MUST_COMPLETE, RequirementType.RECOMMENDED} <= types
    rejected = [r for o in outcomes for r in o.rejected]
    assert any(r["sentence"].startswith("This policy defines") and r["reason"].startswith("informational")
               for r in rejected)
    assert not any(r.statement.startswith("This policy defines") for r in reqs)


def test_extraction_details_for_one_clause():
    outcomes = extract("POL-INFOSEC-001_v1.pdf", "POL-INFOSEC-001", "INFOSEC_POLICY")
    r = next(r for o in outcomes for r in o.requirements if "Information Security Basics" in r.statement)
    assert r.requirement_type == RequirementType.MUST_COMPLETE and r.is_mandatory
    assert r.priority == Priority.HIGH and r.deadline_days == 7
    assert r.assessment_required and r.source_section_id == "2.3"
    assert r.source_key == "POL-INFOSEC-001:2.3:1"
    assert "modal:must_complete" in r.extraction_method and 0 < r.extraction_confidence <= 1


def test_conditions_and_exceptions_are_recorded():
    outcomes = extract("POL-INFOSEC-001_v1.pdf", "POL-INFOSEC-001", "INFOSEC_POLICY")
    exception = next(r for o in outcomes for r in o.requirements if r.statement.startswith("Exception:"))
    assert exception.requirement_type == RequirementType.OPTIONAL
    assert exception.conditions.get("exception") is True and exception.conditions.get("marker") == "during"


def test_event_deadline_is_not_an_onboarding_deadline():
    outcomes = extract("POL-INFOSEC-001_v1.pdf", "POL-INFOSEC-001", "INFOSEC_POLICY")
    incident = next(r for o in outcomes for r in o.requirements if "within one hour of discovery" in r.statement)
    assert incident.deadline_days is None and incident.priority == Priority.CRITICAL


def test_suspicious_chunk_is_never_read_for_requirements():
    outcomes = extract("FAQ-ENG-002_v1.pdf", "FAQ-ENG-002", "FAQ")
    assert any("suspicious" in w for o in outcomes for w in o.warnings)
    assert not any("ignore all previous" in r.statement.lower() for o in outcomes for r in o.requirements)


@pytest.mark.parametrize("text, days", [
    ("within seven calendar days of their joining date", 7), ("within one working day", 1),
    ("a full two weeks from your joining date", 14), ("within thirty minutes of discovery", 1),
    ("between 10:00 and 16:00 on a working day", None),
])
def test_deadline_parsing(text, days):
    assert deadline_days(text) == days


def test_duration_normalisation_to_hours():
    assert durations("within twenty-four hours")[0]["hours"] == 24
    assert durations("thirty minutes")[0]["hours"] == 0.5


# ------------------------------------------------------------ role mapping
ROLES = [RoleInput(i, code, title, dept, name, dcode) for i, (code, title, dept, name, dcode) in enumerate([
    ("SW_INTERN", "Software Intern", 1, "Engineering", "ENG"), ("BACKEND_DEV", "Backend Developer", 1, "Engineering", "ENG"),
    ("FRONTEND_DEV", "Frontend Developer", 1, "Engineering", "ENG"), ("DEVOPS_ENG", "DevOps Engineer", 2, "Infrastructure", "INFRA"),
    ("QA_ENG", "QA Engineer", 3, "Quality Assurance", "QA"), ("PROJECT_MGR", "Project Manager", 4, "Delivery", "DELIVERY"),
    ("SEO_SPEC", "SEO Specialist", 5, "Marketing", "MKT"), ("UIUX_DESIGNER", "UI/UX Designer", 6, "Design", "DESIGN"),
    ("DATA_ANALYST", "Data Analyst", 7, "Data", "DATA"), ("SUPPORT_ENG", "Technical Support Engineer", 8, "Support", "SUPPORT"),
])]


def mapper(roles=ROLES):
    return RoleMapper(roles, load_yaml("requirement_rules"))


def req(statement, dept=None, rtype=RequirementType.MUST_KNOW, clause=""):
    return RequirementInput("r", "R1", statement, rtype, True, Priority.MEDIUM, dept, clause)


def test_explicit_role_mention_maps_only_that_role():
    mappings = mapper().map(req("Backend Developers must not commit secrets to any repository."))
    titles = {ROLES[m.job_role_id].title for m in mappings}
    assert titles == {"Backend Developer"}
    assert mappings[0].mapping_method == MappingMethod.EXPLICIT_ROLE_MENTION
    assert "backend developers" in mappings[0].mapping_reason.lower()
    assert mappings[0].priority_for_role == Priority.HIGH  # explicit mention boost


def test_all_employees_maps_to_all_ten_roles():
    mappings = mapper().map(req("All employees must enable multi-factor authentication.",
                                rtype=RequirementType.MUST_COMPLETE))
    assert len(mappings) == 10 and {m.mapping_method for m in mappings} == {MappingMethod.APPLIES_TO_ALL}


def test_department_document_maps_to_department_roles():
    mappings = mapper().map(req("Standard releases must be scheduled on a working day.", dept=2))
    assert [ROLES[m.job_role_id].code for m in mappings] == ["DEVOPS_ENG"]
    assert mappings[0].mapping_method == MappingMethod.DEPARTMENT_MATCH


def test_department_named_in_sentence():
    mappings = mapper().map(req("All engineering staff must complete the Secure Coding assessment."))
    assert {ROLES[m.job_role_id].code for m in mappings} == {"SW_INTERN", "BACKEND_DEV", "FRONTEND_DEV"}


def test_except_clause_excludes_the_named_role():
    mappings = mapper().map(req("All employees must join the on-call rota except Software Interns."))
    assert "SW_INTERN" not in {ROLES[m.job_role_id].code for m in mappings}
    assert len(mappings) == 9


def test_continuation_sentence_inherits_clause_role():
    clause = "Exception: a DevOps Engineer may query raw records. The query must be logged within one working day."
    mappings = mapper().map(req("The query must be logged within one working day.", clause=clause))
    assert [ROLES[m.job_role_id].code for m in mappings] == ["DEVOPS_ENG"]


def test_new_role_created_later_is_mapped_without_code_change():
    roles = ROLES + [RoleInput(99, "ML_ENG", "Machine Learning Engineer", 1, "Engineering", "ENG")]
    mappings = RoleMapper(roles, load_yaml("requirement_rules")).map(
        req("Machine Learning Engineers must document every training dataset."))
    assert [m.job_role_id for m in mappings] == [99]
    everyone = RoleMapper(roles, load_yaml("requirement_rules")).map(req("All employees must wear a badge."))
    assert 99 in {m.job_role_id for m in everyone}


def test_not_applicable_is_never_mapped():
    assert mapper().map(req("This clause does not apply to contractors.", rtype=RequirementType.NOT_APPLICABLE)) == []


# ------------------------------------------------------------ prerequisites
def test_prerequisite_from_named_artefact():
    nodes = [ReqNode("a", "R1", "DevOps Engineers must demonstrate completion of the Production Access "
                                "Certification before receiving standing production credentials.", "Production Access"),
             ReqNode("b", "R2", "The DevOps Engineer must hold a valid Production Access Certification before "
                                "running a release.", "Release Management")]
    edges, cycles = detect(nodes, load_yaml("requirement_rules"))
    assert [(e.requirement_code, e.depends_on_code) for e in edges] == [("R2", "R1")]
    assert not cycles


def test_cycle_is_rejected_and_reported():
    rules = load_yaml("requirement_rules")
    nodes = [ReqNode("a", "R1", "Staff must complete the Alpha Training once they finish the Beta Course.", "X"),
             ReqNode("b", "R2", "Staff must complete the Beta Course once they finish the Alpha Training.", "X")]
    edges, cycles = detect(nodes, rules)
    assert len(edges) == 1 and len(cycles) == 1
    assert set(cycles[0]["cycle"]) == {"R1", "R2"}
