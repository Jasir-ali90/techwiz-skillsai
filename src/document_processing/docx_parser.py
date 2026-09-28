from docx import Document as DocxDocument

from src.document_processing.models import ParsedDocument, RawBlock


def parse_docx(path: str) -> ParsedDocument:
    """Extracts paragraphs and table text, keeping the paragraph index.
    SRS Step 6: paragraph or section reference is required for DOCX sources."""
    parsed = ParsedDocument(parser="python-docx")
    document = DocxDocument(path)

    index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        index += 1
        style = (paragraph.style.name if paragraph.style else "") or ""
        parsed.blocks.append(
            RawBlock(
                text=text,
                paragraph_index=index,
                is_heading=style.lower().startswith("heading") or style.lower() == "title",
                style_name=style,
            )
        )

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if not cells:
                continue
            index += 1
            parsed.blocks.append(
                RawBlock(text=" | ".join(cells), paragraph_index=index, style_name="Table")
            )

    if not parsed.blocks:
        parsed.warnings.append("The document contained no readable paragraphs or tables")
    return parsed